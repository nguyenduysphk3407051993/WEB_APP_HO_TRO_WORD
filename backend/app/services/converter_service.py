"""Chuyen doi LaTeX <-> MathML <-> OMML va Markdown -> DOCX."""
from __future__ import annotations

from copy import deepcopy
import logging
import re
import unicodedata
from pathlib import Path

import pypandoc
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph

logger = logging.getLogger(__name__)

_MATHML_TAG_RE = re.compile(r"<math\b[^>]*>.*?</math>", re.DOTALL | re.IGNORECASE)
_MARKDOWN_DOCX_FORMAT = (
    "markdown+raw_tex+tex_math_dollars+pipe_tables+grid_tables"
)
_MARKDOWN_TOGGLE_TEX_DOCX_FORMAT = (
    "markdown-raw_tex-tex_math_dollars-tex_math_single_backslash-"
    "tex_math_double_backslash+pipe_tables+grid_tables"
)
_DISPLAY_BRACKET_RE = re.compile(r"\\\[(.+?)\\\]", re.DOTALL)
_INLINE_PAREN_RE = re.compile(r"\\\((.+?)\\\)", re.DOTALL)
_DOLLAR_BLOCK_RE = re.compile(r"\$\$(.+?)\$\$", re.DOTALL)
_DOLLAR_INLINE_RE = re.compile(
    r"(?<!\$)\$(?!\$)(.+?)(?<!\$)\$(?!\$)",
    re.DOTALL,
)
_QUESTION_PREFIX_RE = re.compile(
    r"^\s*((?:Câu|Bài|Ví dụ)\s+\d+[A-Za-z]?(?:\s*[:.)-])?)",
    re.IGNORECASE,
)
_OPTION_PREFIX_RE = re.compile(r"^\s*([A-D])([.)])\s*(.+?)\s*$")
_SUBITEM_PREFIX_RE = re.compile(r"^\s*([a-d])([.)])\s*(.+?)\s*$")
_INLINE_STRUCTURE_RE = re.compile(
    r"(?<!\S)(?P<option>[A-D][.)])\s+|(?<!\S)(?P<subitem>[a-d][.)])\s+"
)
_UNICODE_TEX_REPLACEMENTS = {
    "≤": r"\le ",
    "≥": r"\ge ",
    "≠": r"\ne ",
    "≈": r"\approx ",
    "×": r"\times ",
    "÷": r"\div ",
    "·": r"\cdot ",
    "±": r"\pm ",
    "∞": r"\infty ",
    "∈": r"\in ",
    "∉": r"\notin ",
    "⊂": r"\subset ",
    "⊆": r"\subseteq ",
    "∪": r"\cup ",
    "∩": r"\cap ",
    "∀": r"\forall ",
    "∃": r"\exists ",
    "→": r"\to ",
    "⇒": r"\Rightarrow ",
    "⇔": r"\Leftrightarrow ",
    "∑": r"\sum ",
    "∫": r"\int ",
    "√": r"\sqrt ",
    "π": r"\pi ",
    "α": r"\alpha ",
    "β": r"\beta ",
    "γ": r"\gamma ",
    "δ": r"\delta ",
    "Δ": r"\Delta ",
    "θ": r"\theta ",
    "λ": r"\lambda ",
    "μ": r"\mu ",
    "Ω": r"\Omega ",
    "–": "-",
    "—": "-",
    "−": "-",
}

# Màu tiền tố câu hỏi và marker phương án: #0000FF
_QUESTION_PREFIX_COLOR = RGBColor(0x00, 0x00, 0xFF)
_OPTION_PREFIX_COLOR   = RGBColor(0x00, 0x00, 0xFF)

# Màu từ khoá lời giải: đỏ đậm #C60C4A
_SOLUTION_KEYWORD_COLOR = RGBColor(0xC6, 0x0C, 0x4A)
_SOLUTION_KEYWORDS_RE   = re.compile(
    r"(Lời giải\s*[:.]|Hướng dẫn giải\s*[:.]|Đáp án\s*[:.:]?)",
    re.IGNORECASE,
)

_DOCUMENT_FONT_NAME = "Times New Roman"
_DOCUMENT_BODY_FONT_SIZE = 12
_DOCUMENT_PAGE_WIDTH = Cm(21.59)
_DOCUMENT_PAGE_HEIGHT = Cm(27.94)
_DOCUMENT_LEFT_MARGIN = Cm(2.0)
_DOCUMENT_RIGHT_MARGIN = Cm(1.25)
_DOCUMENT_TOP_MARGIN = Cm(1.5)
_DOCUMENT_BOTTOM_MARGIN = Cm(1.0)
_OPTION_LEFT_INDENT = Cm(0.5)
_OPTION_TWO_COLUMN_TABS = (Cm(0.5), Cm(9.0))
_OPTION_FOUR_COLUMN_TABS = (Cm(0.5), Cm(4.5), Cm(9.0), Cm(13.0))
_OPTION_FOUR_COLUMN_MAX_LENGTH = 20
_OPTION_FOUR_COLUMN_TOTAL_LENGTH = 80
_OPTION_TWO_COLUMN_MAX_LENGTH = 40


class ConverterService:
    """Bo chuyen doi 2 chieu cho LaTeX / MathML / OMML."""

    def latex_to_mathml(self, latex: str, display: bool = True) -> str:
        """Tra ve chuoi MathML (<math>) de dan vao MathType."""
        wrapper = f"\\[\n{latex}\n\\]" if display else f"${latex}$"
        html = pypandoc.convert_text(
            wrapper,
            "html",
            format="latex",
            extra_args=["--mathml"],
        )
        match = _MATHML_TAG_RE.search(html)
        if not match:
            raise ValueError("Khong trich xuat duoc MathML tu LaTeX da cho.")
        mathml = match.group(0)
        if 'xmlns="http://www.w3.org/1998/Math/MathML"' not in mathml:
            mathml = mathml.replace(
                "<math",
                '<math xmlns="http://www.w3.org/1998/Math/MathML"',
                1,
            )
        return mathml

    def mathml_to_latex(self, mathml: str) -> str:
        """MathML -> LaTeX (di qua HTML wrapper)."""
        html = (
            "<!DOCTYPE html><html><head><meta charset='utf-8'></head>"
            f"<body><p>{mathml}</p></body></html>"
        )
        latex = pypandoc.convert_text(html, "latex", format="html")
        latex = latex.strip()
        latex = re.sub(r"^\\\[", "", latex)
        latex = re.sub(r"\\\]$", "", latex)
        latex = re.sub(r"^\$+", "", latex)
        latex = re.sub(r"\$+$", "", latex)
        return latex.strip()

    def latex_to_docx(
        self,
        latex_content: str,
        output_path: Path,
        as_full_document: bool = False,
    ) -> Path:
        """Chuyen LaTeX sang file .docx co cong thuc OMML."""
        if not as_full_document:
            latex_doc = (
                "\\documentclass{article}\n"
                "\\usepackage{amsmath,amssymb,amsfonts}\n"
                "\\begin{document}\n"
                f"{latex_content}\n"
                "\\end{document}\n"
            )
        else:
            latex_doc = latex_content

        output_path.parent.mkdir(parents=True, exist_ok=True)
        pypandoc.convert_text(
            latex_doc,
            "docx",
            format="latex",
            outputfile=str(output_path),
            extra_args=["--mathml"],
        )
        return output_path

    def markdown_to_docx(
        self,
        markdown_content: str,
        output_path: Path,
        resource_path: Path | None = None,
    ) -> Path:
        """Chuyen Markdown co cong thuc $...$/$$...$$ sang DOCX."""
        payload = markdown_content.strip() or " "
        output_path.parent.mkdir(parents=True, exist_ok=True)
        extra_args = ["--wrap=preserve"]
        if resource_path:
            extra_args.append(f"--resource-path={resource_path}")
        pypandoc.convert_text(
            payload,
            "docx",
            format=_MARKDOWN_DOCX_FORMAT,
            outputfile=str(output_path),
            extra_args=extra_args,
        )
        self.postprocess_ocr_docx(output_path)
        return output_path

    def markdown_to_toggle_tex_docx(
        self,
        markdown_content: str,
        output_path: Path,
        resource_path: Path | None = None,
    ) -> Path:
        """Chuyen Markdown sang DOCX, giu cong thuc o dang TeX text cho MathType Toggle TeX."""
        payload = self._normalize_toggle_tex_markdown(markdown_content.strip() or " ")
        output_path.parent.mkdir(parents=True, exist_ok=True)
        extra_args = ["--wrap=preserve"]
        if resource_path:
            extra_args.append(f"--resource-path={resource_path}")
        pypandoc.convert_text(
            payload,
            "docx",
            format=_MARKDOWN_TOGGLE_TEX_DOCX_FORMAT,
            outputfile=str(output_path),
            extra_args=extra_args,
        )
        self.postprocess_ocr_docx(output_path)
        return output_path

    @staticmethod
    def _normalize_toggle_tex_markdown(markdown_content: str) -> str:
        """Dua delimiter ve $...$/\\[...\\] va lam sach TeX cho MathType Toggle TeX."""
        markdown_content = _DOLLAR_BLOCK_RE.sub(
            lambda match: f"\\[\n{ConverterService.sanitize_toggle_tex_latex(match.group(1))}\n\\]",
            markdown_content,
        )
        markdown_content = _DISPLAY_BRACKET_RE.sub(
            lambda match: f"\\[\n{ConverterService.sanitize_toggle_tex_latex(match.group(1))}\n\\]",
            markdown_content,
        )
        markdown_content = _INLINE_PAREN_RE.sub(
            lambda match: f"${ConverterService.sanitize_toggle_tex_latex(match.group(1))}$",
            markdown_content,
        )
        markdown_content = _DOLLAR_INLINE_RE.sub(
            lambda match: f"${ConverterService.sanitize_toggle_tex_latex(match.group(1))}$",
            markdown_content,
        )
        return markdown_content

    @staticmethod
    def sanitize_toggle_tex_latex(latex: str) -> str:
        """Loai Unicode/tieng Viet trong vung TeX de MathType Toggle TeX on dinh hon."""
        latex = latex.strip()
        for source, replacement in _UNICODE_TEX_REPLACEMENTS.items():
            latex = latex.replace(source, replacement)

        latex = latex.replace("đ", "d").replace("Đ", "D")
        latex = unicodedata.normalize("NFKD", latex)
        latex = "".join(char for char in latex if not unicodedata.combining(char))
        latex = latex.encode("ascii", "ignore").decode("ascii")
        latex = re.sub(r"[ \t]+", " ", latex)
        latex = re.sub(r"\s*\n\s*", "\n", latex)
        return latex.strip()

    def postprocess_ocr_docx(self, docx_path: Path) -> Path:
        """Hau xu ly DOCX OCR de dinh dang cau hoi, phuong an, bang, anh."""
        document = Document(docx_path)
        self._apply_document_geometry(document)
        self._apply_document_font(document)
        original_paragraphs = list(document.paragraphs)

        for paragraph in original_paragraphs:
            if not paragraph.text or not paragraph.text.strip():
                continue

            segments = self._explode_structured_paragraph(paragraph.text)
            current = paragraph

            if len(segments) > 1:
                self._set_paragraph_text(current, segments[0][1])
                self._apply_segment_format(current, segments[0][0])

                for kind, text in segments[1:]:
                    current = self._insert_paragraph_after(current, text)
                    self._apply_segment_format(current, kind)
                continue

            self._set_paragraph_text(current, segments[0][1])
            self._apply_segment_format(current, segments[0][0])

        self._layout_answer_choices(document)
        self._split_solution_keyword_paragraphs(document)
        self._split_solution_content_paragraphs(document)
        self._format_solution_sections(document)
        self._format_tables(document)
        self._format_images(document)

        for paragraph in document.paragraphs:
            self._highlight_solution_keywords(paragraph)

        document.save(docx_path)
        return docx_path

    # ──────────────────────────────────────────────────────────────
    # Page geometry
    # ──────────────────────────────────────────────────────────────

    @staticmethod
    def _apply_document_geometry(document: Document) -> None:
        for section in document.sections:
            section.page_width = _DOCUMENT_PAGE_WIDTH
            section.page_height = _DOCUMENT_PAGE_HEIGHT
            section.left_margin = _DOCUMENT_LEFT_MARGIN
            section.right_margin = _DOCUMENT_RIGHT_MARGIN
            section.top_margin = _DOCUMENT_TOP_MARGIN
            section.bottom_margin = _DOCUMENT_BOTTOM_MARGIN

    # ──────────────────────────────────────────────────────────────
    # Font — Times New Roman toàn bộ
    # ──────────────────────────────────────────────────────────────

    def _apply_document_font(self, document: Document) -> None:
        for style in document.styles:
            if style.type != WD_STYLE_TYPE.PARAGRAPH:
                continue
            self._set_font_family(style.font, style.element.get_or_add_rPr())

        normal_style = document.styles["Normal"]
        normal_style.font.size = Pt(_DOCUMENT_BODY_FONT_SIZE)

        for paragraph in self._iter_document_paragraphs(document):
            for run in paragraph.runs:
                self._set_run_font(run, size=None)

    @staticmethod
    def _iter_document_paragraphs(parent):
        for paragraph in parent.paragraphs:
            yield paragraph
        for table in parent.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from ConverterService._iter_document_paragraphs(cell)

    # ──────────────────────────────────────────────────────────────
    # Layout A–D answer choices (4-col / 2-col / 1-per-line)
    # ──────────────────────────────────────────────────────────────

    def _layout_answer_choices(self, document: Document) -> None:
        paragraphs = list(document.paragraphs)
        index = 0

        while index < len(paragraphs):
            first = self._parse_choice(paragraphs[index].text)
            if not first or first[0] != "A":
                index += 1
                continue

            group: list[tuple[Paragraph, str, str]] = []
            expected_markers = ("A", "B", "C", "D")
            for offset, expected in enumerate(expected_markers):
                if index + offset >= len(paragraphs):
                    break
                parsed = self._parse_choice(paragraphs[index + offset].text)
                if not parsed or parsed[0] != expected:
                    break
                group.append((paragraphs[index + offset], parsed[0], parsed[1]))

            if len(group) != 4:
                index += 1
                continue

            choices = [(marker, content) for _, marker, content in group]
            first_element = group[0][0]._p
            if self._use_four_choice_columns(choices):
                self._set_choice_row(group[0][0], choices, _OPTION_FOUR_COLUMN_TABS)
                for paragraph, _, _ in group[1:]:
                    self._remove_paragraph(paragraph)
            elif self._use_two_choice_columns(choices):
                self._set_choice_row(group[0][0], choices[:2], _OPTION_TWO_COLUMN_TABS)
                self._set_choice_row(group[2][0], choices[2:], _OPTION_TWO_COLUMN_TABS)
                self._remove_paragraph(group[1][0])
                self._remove_paragraph(group[3][0])
            else:
                # 1 phương án/dòng — đã format bởi _format_choice_paragraph
                index += 1
                continue

            paragraphs = list(document.paragraphs)
            index = next(
                (i for i, p in enumerate(paragraphs) if p._p is first_element),
                index,
            ) + 1

    @staticmethod
    def _parse_choice(text: str) -> tuple[str, str] | None:
        match = _OPTION_PREFIX_RE.match((text or "").strip())
        if not match:
            return None
        return match.group(1), match.group(3).strip()

    @staticmethod
    def _use_four_choice_columns(choices: list[tuple[str, str]]) -> bool:
        lengths = [len(content) for _, content in choices]
        return (
            max(lengths, default=0) <= _OPTION_FOUR_COLUMN_MAX_LENGTH
            and sum(lengths) <= _OPTION_FOUR_COLUMN_TOTAL_LENGTH
        )

    @staticmethod
    def _use_two_choice_columns(choices: list[tuple[str, str]]) -> bool:
        lengths = [len(content) for _, content in choices]
        return max(lengths, default=0) <= _OPTION_TWO_COLUMN_MAX_LENGTH

    def _set_choice_row(
        self,
        paragraph: Paragraph,
        choices: list[tuple[str, str]],
        tab_positions: tuple,
    ) -> None:
        self._clear_paragraph(paragraph)
        paragraph.style = "Normal"
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        fmt = paragraph.paragraph_format
        fmt.left_indent = _OPTION_LEFT_INDENT
        fmt.first_line_indent = Pt(0)
        fmt.space_before = Pt(3)
        fmt.space_after = Pt(3)
        fmt.line_spacing = 1
        self._reset_tab_stops(paragraph)
        for position in tab_positions:
            fmt.tab_stops.add_tab_stop(position, WD_TAB_ALIGNMENT.LEFT, WD_TAB_LEADER.SPACES)

        for choice_index, (marker, content) in enumerate(choices):
            if choice_index:
                tab_run = paragraph.add_run("\t")
                self._set_run_font(tab_run)
            marker_run = paragraph.add_run(f"{marker}. ")
            self._set_run_font(marker_run, bold=True)
            marker_run.font.color.rgb = _OPTION_PREFIX_COLOR
            content_run = paragraph.add_run(content)
            self._set_run_font(content_run, bold=False)

    # ──────────────────────────────────────────────────────────────
    # Paragraph structure: explode inline options/subitems
    # ──────────────────────────────────────────────────────────────

    @staticmethod
    def _explode_structured_paragraph(text: str) -> list[tuple[str, str]]:
        text = text.replace("\xa0", " ")
        text = re.sub(r"[ \t]+", " ", text).strip()
        if not text:
            return [("paragraph", "")]

        segments: list[tuple[str, str]] = []
        question_match = _QUESTION_PREFIX_RE.match(text)

        if question_match:
            first_marker = _INLINE_STRUCTURE_RE.search(text, question_match.end())
            if not first_marker:
                return [("question", text)]

            question_text = text[: first_marker.start()].strip()
            if question_text:
                segments.append(("question", question_text))
            text = text[first_marker.start():].strip()

        markers = list(_INLINE_STRUCTURE_RE.finditer(text))
        if not markers:
            if segments:
                return segments
            if _OPTION_PREFIX_RE.match(text):
                return [("option", text)]
            if _SUBITEM_PREFIX_RE.match(text):
                return [("subitem", text)]
            return [("paragraph", text)]

        if not segments and markers[0].start() > 0:
            prefix_text = text[: markers[0].start()].strip()
            if prefix_text:
                segments.append(("paragraph", prefix_text))

        for index, marker in enumerate(markers):
            end = markers[index + 1].start() if index + 1 < len(markers) else len(text)
            chunk = text[marker.start(): end].strip()
            if not chunk:
                continue
            kind = "option" if marker.group("option") else "subitem"
            segments.append((kind, chunk))

        return segments or [("paragraph", text)]

    def _apply_segment_format(self, paragraph: Paragraph, kind: str) -> None:
        if kind == "question":
            self._format_question_prefix(paragraph)
        elif kind == "option":
            self._format_choice_paragraph(paragraph, is_subitem=False)
        elif kind == "subitem":
            self._format_choice_paragraph(paragraph, is_subitem=True)

    def _format_question_prefix(self, paragraph: Paragraph) -> None:
        text = paragraph.text.strip()
        match = _QUESTION_PREFIX_RE.match(text)
        if not match:
            return

        prefix = match.group(1).strip()
        remainder = text[match.end():].lstrip()
        self._clear_paragraph(paragraph)

        prefix_run = paragraph.add_run(prefix)
        prefix_run.bold = True
        prefix_run.font.color.rgb = _QUESTION_PREFIX_COLOR

        if remainder:
            paragraph.add_run(f" {remainder}")

    def _format_choice_paragraph(self, paragraph: Paragraph, *, is_subitem: bool) -> None:
        text = paragraph.text.strip()
        pattern = _SUBITEM_PREFIX_RE if is_subitem else _OPTION_PREFIX_RE
        match = pattern.match(text)
        if not match:
            return

        marker = "".join(match.groups()[:-1])
        content = match.group(len(match.groups())).strip()
        self._clear_paragraph(paragraph)

        marker_run = paragraph.add_run(f"{marker} ")
        content_run = paragraph.add_run(content)

        fmt = paragraph.paragraph_format
        fmt.left_indent = _OPTION_LEFT_INDENT
        fmt.first_line_indent = Pt(0)
        fmt.space_before = Pt(3)
        fmt.space_after = Pt(3)
        fmt.line_spacing = 1
        self._reset_tab_stops(paragraph)
        fmt.tab_stops.add_tab_stop(_OPTION_LEFT_INDENT, WD_TAB_ALIGNMENT.LEFT, WD_TAB_LEADER.SPACES)
        self._set_run_font(marker_run, bold=not is_subitem)
        if not is_subitem:
            marker_run.font.color.rgb = _OPTION_PREFIX_COLOR
        self._set_run_font(content_run, bold=False)

    # ──────────────────────────────────────────────────────────────
    # Tách keyword lời giải xuống dòng riêng
    # ──────────────────────────────────────────────────────────────

    def _split_solution_keyword_paragraphs(self, document: Document) -> None:
        """Nếu paragraph BẮT ĐẦU bằng 'Lời giải:/Đáp án:...' + nội dung, tách thành 2 đoạn."""
        for paragraph in list(document.paragraphs):
            text = paragraph.text.strip()
            if not text:
                continue
            # Chỉ match tại đầu đoạn — tránh phá nội dung câu hỏi có chứa "đáp án" giữa câu
            match = _SOLUTION_KEYWORDS_RE.match(text)
            if not match:
                continue
            after = text[match.end():].strip()
            if not after:
                continue
            # Giữ keyword ở đoạn hiện tại, nội dung còn lại sang đoạn mới bên dưới
            self._clear_paragraph(paragraph)
            paragraph.add_run(match.group(1))
            self._insert_paragraph_after(paragraph, after)

    # ──────────────────────────────────────────────────────────────
    # Highlight "Lời giải", "Hướng dẫn giải", "Đáp án" in đậm đỏ
    # ──────────────────────────────────────────────────────────────

    def _highlight_solution_keywords(self, paragraph: Paragraph) -> None:
        runs = list(paragraph.runs)
        for run in runs:
            text = run.text
            if not text:
                continue

            matches = list(_SOLUTION_KEYWORDS_RE.finditer(text))
            if not matches:
                continue

            run_element = run._element
            last_end = 0

            for match in matches:
                before_text = text[last_end: match.start()]
                if before_text:
                    new_run = paragraph.add_run(before_text)
                    new_run.bold = run.bold
                    new_run.italic = run.italic
                    self._set_run_font(new_run)
                    run_element.addprevious(new_run._element)

                kw_run = paragraph.add_run(match.group(1))
                kw_run.bold = True
                kw_run.font.color.rgb = _SOLUTION_KEYWORD_COLOR
                self._set_run_font(kw_run)
                run_element.addprevious(kw_run._element)

                last_end = match.end()

            after_text = text[last_end:]
            if after_text:
                new_run = paragraph.add_run(after_text)
                new_run.bold = run.bold
                new_run.italic = run.italic
                self._set_run_font(new_run)
                run_element.addprevious(new_run._element)

            paragraph._p.remove(run_element)

    # ──────────────────────────────────────────────────────────────
    # Tách ý nhúng trong 1 đoạn lời giải thành các đoạn riêng
    # ──────────────────────────────────────────────────────────────

    # Pattern nhận diện ký tự bắt đầu ý mới (chỉ dùng để tách — không ảnh hưởng math)
    _INLINE_ITEM_SEP_RE = re.compile(
        r"(?:^|\s{2,}|(?<=[.!?:])[ \t]+)"          # bắt đầu dòng hoặc sau kết thúc câu
        r"([-*•–]|[+]|=>|⇒|\d{1,2}[.)]|[a-d][.)])" # marker
        r"(?=\s)",
        re.MULTILINE,
    )

    def _split_solution_content_paragraphs(self, document: Document) -> None:
        """Tách ý nhúng trong cùng 1 paragraph lời giải thành các đoạn riêng."""
        in_solution = False
        for paragraph in list(document.paragraphs):
            text = paragraph.text.strip()
            if not text:
                continue
            if _QUESTION_PREFIX_RE.match(text):
                in_solution = False
                continue
            if _SOLUTION_KEYWORDS_RE.match(text):
                in_solution = True
                continue
            if not in_solution:
                continue

            # Tìm tất cả marker bắt đầu ý mới bên trong text
            matches = list(self._INLINE_ITEM_SEP_RE.finditer(text))
            # Bỏ match ở vị trí 0 (đầu dòng — đây là marker của chính paragraph này)
            inner = [m for m in matches if m.start() > 0]
            if not inner:
                continue

            # Tách text tại vị trí các marker nhúng
            current_para = paragraph
            prev_end = 0
            splits: list[str] = []
            for m in inner:
                splits.append(text[prev_end: m.start()].strip())
                prev_end = m.start()
            splits.append(text[prev_end:].strip())

            if len(splits) < 2:
                continue

            # Ghi lại đoạn đầu, chèn các đoạn tiếp theo bên dưới
            self._clear_paragraph(current_para)
            current_para.add_run(splits[0])
            for chunk in reversed(splits[1:]):
                new_para = self._insert_paragraph_after(current_para, chunk)
                _ = new_para  # paragraphs sẽ được format bởi _format_solution_sections

    # ──────────────────────────────────────────────────────────────
    # Format solution sections — phân cấp chính/phụ rõ ràng
    # ──────────────────────────────────────────────────────────────

    def _format_solution_sections(self, document: Document) -> None:
        in_solution = False
        # Level 1: dấu đầu dòng chính
        l1_re = re.compile(
            r"^([-*•–—]|[a-dA-D][.)]|\d{1,2}[.)])\s+(.+)$", re.DOTALL
        )
        # Level 2: dấu phụ
        l2_re = re.compile(r"^([+>]|\+\+)\s+(.+)$", re.DOTALL)
        # Kết luận
        concl_re = re.compile(
            r"^(=>|⇒|Vậy\b[,:]?|Kết luận\b[,:]?|Chọn\b[,:]?)\s*(.*)$",
            re.IGNORECASE,
        )

        for paragraph in list(document.paragraphs):
            text = paragraph.text.strip()
            if not text:
                continue

            if _QUESTION_PREFIX_RE.match(text) or text.startswith("#"):
                in_solution = False
                continue

            if _SOLUTION_KEYWORDS_RE.match(text):
                in_solution = True
                fmt = paragraph.paragraph_format
                fmt.space_before = Pt(6)
                fmt.space_after = Pt(2)
                fmt.left_indent = Pt(0)
                fmt.first_line_indent = Pt(0)
                fmt.line_spacing = 1.15
                fmt.keep_with_next = True
                continue

            if not in_solution:
                continue

            fmt = paragraph.paragraph_format
            fmt.line_spacing = 1.15

            # Kiểm tra list natively từ pandoc (numPr / List style)
            p_pr = paragraph._p.get_or_add_pPr()
            num_pr = p_pr.find(qn("w:numPr"))
            style_name = paragraph.style.name
            is_native_list = (num_pr is not None) or style_name.startswith("List")

            if is_native_list:
                ilvl = 0
                if num_pr is not None:
                    ilvl_elem = num_pr.find(qn("w:ilvl"))
                    if ilvl_elem is not None:
                        try:
                            ilvl = int(ilvl_elem.get(qn("w:val"), "0"))
                        except ValueError:
                            pass
                else:
                    if any(d in style_name for d in ("2", "3", "4")):
                        ilvl = 1
                if ilvl >= 1:
                    # L2 phụ — block indent, first line thẳng hàng với continuation
                    fmt.left_indent = Cm(1.0)
                    fmt.first_line_indent = Pt(0)
                    fmt.space_before = Pt(2)
                    fmt.space_after = Pt(2)
                else:
                    # L1 chính — block indent
                    fmt.left_indent = Cm(0.5)
                    fmt.first_line_indent = Pt(0)
                    fmt.space_before = Pt(5)
                    fmt.space_after = Pt(2)
                continue

            l2_match = l2_re.match(text)
            l1_match = l1_re.match(text)
            concl_match = concl_re.match(text)

            if l2_match:
                # L2 phụ — block indent sâu hơn L1, first line = left indent
                fmt.left_indent = Cm(1.0)
                fmt.first_line_indent = Pt(0)
                fmt.space_before = Pt(2)
                fmt.space_after = Pt(2)
                self._reset_tab_stops(paragraph)

            elif l1_match:
                # L1 chính — block indent, marker in đậm, first line = left indent
                fmt.left_indent = Cm(0.5)
                fmt.first_line_indent = Pt(0)
                fmt.space_before = Pt(5)
                fmt.space_after = Pt(2)
                self._reset_tab_stops(paragraph)
                # Bold marker run đầu tiên
                if paragraph.runs:
                    paragraph.runs[0].bold = True

            elif concl_match:
                # Kết luận — tách biệt rõ, in đậm nghiêng
                fmt.left_indent = Cm(0.5)
                fmt.first_line_indent = Pt(0)
                fmt.space_before = Pt(7)
                fmt.space_after = Pt(3)
                for run in paragraph.runs:
                    run.bold = True
                    run.italic = True

            else:
                # Văn xuôi thông thường trong lời giải
                fmt.left_indent = Cm(0.5)
                fmt.first_line_indent = Pt(0)
                fmt.space_before = Pt(3)
                fmt.space_after = Pt(2)

    @staticmethod
    def _convert_manual_prefix_to_tab(paragraph: Paragraph, prefix: str) -> None:
        text_so_far = ""
        for run in paragraph.runs:
            if not run.text:
                continue
            text_so_far += run.text
            if text_so_far.lstrip().startswith(prefix):
                stripped = run.text.lstrip()
                if stripped.startswith(prefix + " "):
                    lead = run.text[: len(run.text) - len(stripped)]
                    run.text = f"{lead}{prefix}\t{stripped[len(prefix) + 1:]}"
                    break
                elif stripped.startswith(prefix) and len(stripped) > len(prefix) and stripped[len(prefix)] in (" ", "\t"):
                    lead = run.text[: len(run.text) - len(stripped)]
                    rest = stripped[len(prefix):].lstrip(" \t")
                    run.text = f"{lead}{prefix}\t{rest}"
                    break

    # ──────────────────────────────────────────────────────────────
    # Tables: header shading, borders, cell padding
    # ──────────────────────────────────────────────────────────────

    def _format_tables(self, document: Document) -> None:
        for table in document.tables:
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            try:
                self._set_table_borders(table)
            except Exception as exc:
                logger.warning("Loi ve vien bang: %s", exc)

            if not table.rows:
                continue

            # Header row
            header_row = table.rows[0]
            tr_pr = header_row._tr.get_or_add_trPr()
            tr_pr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
            tr_pr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
            for cell in header_row.cells:
                self._set_cell_background(cell, "F2F2F2")
                self._set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
                for para in cell.paragraphs:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in para.runs:
                        run.bold = True
                        self._set_run_font(run, size=11)

            # Body rows
            for row in table.rows[1:]:
                tr_pr = row._tr.get_or_add_trPr()
                tr_pr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
                for cell in row.cells:
                    self._set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
                    for para in cell.paragraphs:
                        for run in para.runs:
                            self._set_run_font(run, size=11)

    @staticmethod
    def _set_table_borders(table) -> None:
        tbl_pr = table._tbl.tblPr
        borders = tbl_pr.find(qn("w:tblBorders"))
        if borders is not None:
            tbl_pr.remove(borders)
        tbl_pr.append(parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'  <w:top    w:val="single" w:sz="4" w:space="0" w:color="D3D3D3"/>'
            f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="D3D3D3"/>'
            f'  <w:left   w:val="none"/>'
            f'  <w:right  w:val="none"/>'
            f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/>'
            f'  <w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        ))

    @staticmethod
    def _set_cell_background(cell, hex_color: str) -> None:
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    @staticmethod
    def _set_cell_margins(cell, *, top=100, bottom=100, left=150, right=150) -> None:
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_pr.append(parse_xml(
            f'<w:tcMar {nsdecls("w")}>'
            f'  <w:top    w:w="{top}"    w:type="dxa"/>'
            f'  <w:bottom w:w="{bottom}" w:type="dxa"/>'
            f'  <w:left   w:w="{left}"   w:type="dxa"/>'
            f'  <w:right  w:w="{right}"  w:type="dxa"/>'
            f'</w:tcMar>'
        ))

    # ──────────────────────────────────────────────────────────────
    # Images: căn giữa + giới hạn chiều rộng 12 cm
    # ──────────────────────────────────────────────────────────────

    @staticmethod
    def _format_images(document: Document) -> None:
        for paragraph in document.paragraphs:
            has_drawing = any(
                run._element.find(qn("w:drawing")) is not None
                for run in paragraph.runs
            )
            if has_drawing:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_before = Pt(6)
                paragraph.paragraph_format.space_after = Pt(6)

        max_width = Cm(12.0)
        for shape in document.inline_shapes:
            try:
                if shape.width > max_width:
                    ratio = max_width / shape.width
                    shape.width = max_width
                    shape.height = int(shape.height * ratio)
            except Exception:
                pass

    # ──────────────────────────────────────────────────────────────
    # Paragraph helpers
    # ──────────────────────────────────────────────────────────────

    @staticmethod
    def _set_font_family(font, r_pr) -> None:
        font.name = _DOCUMENT_FONT_NAME
        r_fonts = r_pr.get_or_add_rFonts()
        for name in ("ascii", "hAnsi", "eastAsia", "cs"):
            r_fonts.set(qn(f"w:{name}"), _DOCUMENT_FONT_NAME)

    @classmethod
    def _set_run_font(
        cls,
        run,
        *,
        size: int | None = _DOCUMENT_BODY_FONT_SIZE,
        bold: bool | None = None,
    ) -> None:
        r_pr = run._element.get_or_add_rPr()
        cls._set_font_family(run.font, r_pr)
        if size is not None:
            run.font.size = Pt(size)
        if bold is not None:
            run.bold = bold

    @staticmethod
    def _set_paragraph_text(paragraph: Paragraph, text: str) -> None:
        if paragraph.text == text:
            return
        paragraph.text = text

    @staticmethod
    def _clear_paragraph(paragraph: Paragraph) -> None:
        for child in list(paragraph._p):
            if child.tag == qn("w:pPr"):
                continue
            paragraph._p.remove(child)

    @staticmethod
    def _reset_tab_stops(paragraph: Paragraph) -> None:
        p_pr = paragraph._p.get_or_add_pPr()
        tabs = p_pr.find(qn("w:tabs"))
        if tabs is not None:
            p_pr.remove(tabs)

    @staticmethod
    def _insert_paragraph_after(paragraph: Paragraph, text: str) -> Paragraph:
        new_p = deepcopy(paragraph._p)
        for child in list(new_p):
            if child.tag != qn("w:pPr"):
                new_p.remove(child)
        paragraph._p.addnext(new_p)
        new_paragraph = Paragraph(new_p, paragraph._parent)
        if text:
            new_paragraph.add_run(text)
        return new_paragraph

    @staticmethod
    def _remove_paragraph(paragraph: Paragraph) -> None:
        element = paragraph._element
        element.getparent().remove(element)
        paragraph._p = paragraph._element = None

    # ──────────────────────────────────────────────────────────────
    # Docx → LaTeX / MathType helpers
    # ──────────────────────────────────────────────────────────────

    def docx_to_latex(self, docx_path: Path) -> str:
        """File .docx (chua OMML/Equation) -> LaTeX."""
        latex = pypandoc.convert_file(
            str(docx_path),
            "latex",
            extra_args=["--wrap=preserve"],
        )
        return latex

    def latex_to_mathtype_xml(self, latex: str) -> dict:
        """Chuan bi noi dung de paste vao MathType."""
        mathml = self.latex_to_mathml(latex, display=True)
        copy_payload = "<html><body>" f"{mathml}" "</body></html>"
        return {"mathml": mathml, "html_payload": copy_payload}


converter_service = ConverterService()
